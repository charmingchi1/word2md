#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Word(.docx) / PDF / PPT(.pptx) 转 Markdown 脚本(含图片文字识别)

用法:
    python word2md.py                          # 自动检测脚本所在目录下所有 .docx/.pdf/.pptx
    python word2md.py test.docx report.pdf     # 指定文件, docx/pdf/pptx 可混选
    python word2md.py docs/ -o output/         # 目录批量转换(递归查找)
    python word2md.py test.pdf --no-ocr        # 只提取图片, 不识别图片中的文字

说明:
    - .docx: mammoth 解析为 HTML 再转 Markdown, 标题/加粗/斜体/列表/表格/超链接支持。
    - .pdf: pymupdf4llm 提取文本/表格/图片, 自带标题识别。
    - .pptx: python-pptx 逐页提取, 每页标题作为二级标题, 正文按层级转为列表,
      表格转 GFM 表格, 计算图表数据转表格, 批注(备注)以引用块附在页末。
    - 图片自动提取到 Markdown 旁的 images 目录; 默认用 RapidOCR 识别图片中的
      文字, 以引用块附在对应图片下方, 得到"正文 + 图片文字"的完整内容。
      纯图片扫描页也会以这种方式提取出文字。
    - PDF 每页重复的页眉页脚自动去除; 图片/扫描件内的表格会提取出文字行,
      但无法还原表格结构。
    - 旧版 .doc/.ppt 格式不支持, 请先另存为 .docx/.pptx。
"""

import argparse
import os
import re
import sys
from pathlib import Path

import mammoth
from markdownify import MarkdownConverter

SUPPORTED_EXTS = {".docx", ".pdf", ".pptx"}

# 旧版二进制格式: 无法直接解析, 给出另存为提示
LEGACY_EXTS = {".doc": ".docx", ".ppt": ".pptx", ".pps": ".pptx"}

# 中英文样式名都映射一遍: 中文版 Word 内置样式在 XML 里通常仍存英文,
# 但自定义或部分模板文档会存中文名, 双保险确保标题层级不丢
STYLE_MAP = """
p[style-name='Title'] => h1:fresh
p[style-name='Subtitle'] => h2:fresh
p[style-name='heading 1'] => h1:fresh
p[style-name='heading 2'] => h2:fresh
p[style-name='heading 3'] => h3:fresh
p[style-name='heading 4'] => h4:fresh
p[style-name='heading 5'] => h5:fresh
p[style-name='heading 6'] => h6:fresh
p[style-name='标题'] => h1:fresh
p[style-name='副标题'] => h2:fresh
p[style-name='标题 1'] => h1:fresh
p[style-name='标题 2'] => h2:fresh
p[style-name='标题 3'] => h3:fresh
p[style-name='标题 4'] => h4:fresh
p[style-name='标题 5'] => h5:fresh
p[style-name='标题 6'] => h6:fresh
p[style-name='Quote'] => blockquote
p[style-name='Intense Quote'] => blockquote
p[style-name='引用'] => blockquote
p[style-name='明显引用'] => blockquote
r[style-name='Strong'] => strong
r[style-name='Emphasis'] => em
r[style-name='加粗'] => strong
r[style-name='强调'] => em
"""

# Word 内嵌图片 content-type 到扩展名的映射(emf/wmf 无法被 OCR 和网页显示, 仅原样保存)
CONTENT_TYPE_EXTS = {
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/gif": ".gif",
    "image/bmp": ".bmp",
    "image/tiff": ".tiff",
    "image/x-emf": ".emf",
    "image/x-wmf": ".wmf",
    "image/svg+xml": ".svg",
    "image/webp": ".webp",
}

# 可以送入 OCR 的图片格式
OCR_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp", ".tiff"}

_OCR_ENGINE = None


def get_ocr_engine():
    """懒加载 RapidOCR 引擎(全局单例, 避免每个文档重复加载模型); 未安装时返回 None"""
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
            _OCR_ENGINE = RapidOCR()
        except Exception:
            _OCR_ENGINE = False
    return _OCR_ENGINE or None


def run_ocr(image_path):
    """对单张图片做文字识别, 返回按行拼接的文本; 无文字或失败返回 None"""
    engine = get_ocr_engine()
    if engine is None:
        return None
    try:
        result, _ = engine(str(image_path))
    except Exception:
        return None
    # 每条结果为 [坐标框, 文字, 置信度], 引擎已按阅读顺序排列
    lines = [item[1] for item in (result or []) if len(item) >= 2 and isinstance(item[1], str)]
    return "\n".join(lines).strip() or None


def save_image_bytes(data, images_dir, name, do_ocr, ocr_texts):
    """把一张图片写入 images 目录, 可选做 OCR 并把结果存入 ocr_texts[相对路径];
    返回 Markdown 用的相对路径。文件名由调用方给出, 见各格式的参数命名约定"""
    images_dir.mkdir(parents=True, exist_ok=True)
    with open(images_dir / name, "wb") as f:
        f.write(data)
    rel = f"{images_dir.name}/{name}"
    if os.path.splitext(name)[1].lower() in OCR_EXTS:
        text = run_ocr(images_dir / name) if do_ocr else None
        if text:
            ocr_texts[rel] = text
            print(f"       [OCR] {rel}: 识别出 {len(text.splitlines())} 行文字")
    return rel


def make_image_converter(images_dir, doc_stem, do_ocr, ocr_texts):
    """返回 mammoth 图片转换器: 图片落盘(文件名 {文档名}_{序号}), 可选 OCR"""

    def convert_image(image):
        content_type = (image.content_type or "").split(";")[0].strip().lower()
        ext = CONTENT_TYPE_EXTS.get(content_type, ".png")
        seq = convert_image.seq
        convert_image.seq += 1
        with image.open() as data:
            blob = data.read()
        name = f"{doc_stem}_{seq}{ext}"
        return {"src": save_image_bytes(blob, images_dir, name, do_ocr, ocr_texts)}

    convert_image.seq = 1
    return convert_image


def insert_ocr_text(markdown, ocr_texts):
    """把每张图的 OCR 文字以引用块形式追加到 Markdown 中对应图片引用的下方"""
    for rel, text in ocr_texts.items():
        quoted = "\n".join("> " + line for line in text.splitlines())
        # 结尾补空行, 避免紧随其后的正文被 CommonMark 惰性延续规则并入引用块
        block = f"\n\n> **图片中的文字：**\n{quoted}\n"
        pattern = re.compile(r"(!\[[^\]]*\]\(" + re.escape(rel) + r"\))")
        markdown = pattern.sub(lambda m: m.group(1) + block, markdown, count=1)
    return markdown


def docx_to_markdown(docx_path, out_path, images_dir_name, do_ocr=True):
    """转换单个 .docx 文件, 返回警告信息列表"""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    images_dir = out_path.parent / images_dir_name
    ocr_texts = {}
    image_converter = make_image_converter(images_dir, docx_path.stem, do_ocr, ocr_texts)

    with open(docx_path, "rb") as f:
        result = mammoth.convert_to_html(
            f,
            style_map=STYLE_MAP,
            convert_image=mammoth.images.img_element(image_converter),
        )

    markdown = MarkdownConverter(
        heading_style="ATX",   # 标题用 # 风格
        bullets="-",           # 无序列表用 -
        autolinks=True,        # 裸链接自动加 <>
    ).convert(result.value)

    markdown = insert_ocr_text(markdown, ocr_texts)
    markdown = markdown.strip() + "\n"

    with open(out_path, "w", encoding="utf-8", newline="\n") as f:
        f.write(markdown)

    return [msg.message for msg in result.messages]


def _pdf_reference_lines(pdf_abs):
    """用 PyMuPDF 直接提取全文文本行(去空白), 作为转换是否丢文本的比对基准"""
    import pymupdf
    lines = []
    with pymupdf.open(pdf_abs) as doc:
        for page in doc:
            for ln in page.get_text().splitlines():
                ln = "".join(ln.split())
                if len(ln) >= 6:  # 忽略过短的行, 避免页码/零散符号造成误判
                    lines.append(ln)
    return lines


def _text_loss_ratio(markdown, ref_lines):
    """正文行在转换结果中缺失的比例; 0.0 表示没有丢失"""
    if not ref_lines:
        return 0.0
    body = "".join(markdown.split())
    missing = sum(1 for ln in ref_lines if ln not in body)
    return missing / len(ref_lines)


def pdf_to_markdown(pdf_path, out_path, images_dir_name, do_ocr=True):
    """转换单个 .pdf 文件, 返回警告信息列表"""
    try:
        import pymupdf4llm
    except ImportError:
        raise RuntimeError("未安装 pymupdf4llm, 无法处理 PDF。安装: pip install pymupdf4llm")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    images_dir = out_path.parent / images_dir_name
    # 先把 pdf 路径解析为绝对路径, 之后的 chdir 才不会影响它
    pdf_abs = str(pdf_path.resolve())

    def _convert(layout):
        # 布局引擎: 图片位置准确、自动去页眉页脚; 传统路径参数集不同
        kwargs = dict(write_images=True, image_path=images_dir_name, show_progress=False)
        if layout:
            kwargs.update(use_ocr=False, header=False, footer=False)
        pymupdf4llm.use_layout(layout)
        # image_path 按当前工作目录解析, 临时切到输出目录,
        # 让图片落在输出目录且 markdown 里的引用是相对路径
        old_cwd = os.getcwd()
        os.chdir(out_path.parent)
        try:
            return pymupdf4llm.to_markdown(pdf_abs, **kwargs)
        finally:
            os.chdir(old_cwd)

    ref_lines = _pdf_reference_lines(pdf_abs)

    # 首选布局引擎; 其对含纯图片页的文档有丢正文 bug, 检测到丢失则回退传统路径
    markdown = _convert(layout=True)
    loss = _text_loss_ratio(markdown, ref_lines)
    if ref_lines and loss > 0.3:
        fallback = _convert(layout=False)
        if _text_loss_ratio(fallback, ref_lines) < loss:
            print("       布局引擎丢失正文, 已自动回退到兼容解析模式")
            markdown = fallback

    # 清理本次转换产生但最终未被引用的图片(两种路径各写一套图片)
    if images_dir.is_dir():
        referenced = set(re.findall(r"!\[[^\]]*\]\(([^)]+)\)", markdown))
        for img in images_dir.glob(pdf_path.name + "-*"):
            if f"{images_dir_name}/{img.name}" not in referenced:
                img.unlink(missing_ok=True)

    # 最终结果里引用的图片依次 OCR
    ocr_texts = {}
    if do_ocr and images_dir.is_dir():
        for img in sorted(images_dir.glob(pdf_path.name + "-*")):
            if img.suffix.lower() not in OCR_EXTS:
                continue
            text = run_ocr(img)
            if text:
                rel = f"{images_dir_name}/{img.name}"
                ocr_texts[rel] = text
                print(f"       [OCR] {rel}: 识别出 {len(text.splitlines())} 行文字")

    markdown = insert_ocr_text(markdown, ocr_texts)
    markdown = markdown.strip() + "\n"

    with open(out_path, "w", encoding="utf-8", newline="\n") as f:
        f.write(markdown)

    return []


def _load_pptx():
    """懒加载 python-pptx(未安装时给出明确提示), 返回 (Presentation, MSO_SHAPE_TYPE, PP_PLACEHOLDER)"""
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
    except ImportError:
        raise RuntimeError("未安装 python-pptx, 无法处理 PPTX。安装: pip install python-pptx")
    return Presentation, MSO_SHAPE_TYPE, PP_PLACEHOLDER


def _render_gfm_table(header, rows):
    """把二维单元格转成 GFM 管道表格(首行作表头); 单元格内换行压成空格, 竖线转义"""
    def cell(value):
        text = " ".join(str(value).split()).replace("|", "\\|")
        return text or " "

    lines = ["| " + " | ".join(cell(c) for c in header) + " |",
             "| " + " | ".join("---" for _ in header) + " |"]
    lines += ["| " + " | ".join(cell(c) for c in row) + " |" for row in rows]
    return "\n".join(lines)


def _pptx_table_to_markdown(table):
    """PPT 表格转 GFM 表格; 首行作表头(PPT 表格通常首行就是表头)"""
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    # 合并单元格会让各行长度不一致, 补齐到最大列数
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    return _render_gfm_table(rows[0], rows[1:])


def _pptx_chart_to_markdown(chart):
    """图表转成标题 + GFM 表格; 数据读不出来时退化成一行说明"""
    def number(value):
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))  # 图表数值取整, 避免 900.0 这样的显示
        return str(value)

    title = ""
    try:
        if chart.has_title:
            title = chart.chart_title.text_frame.text.strip()
        plot = chart.plots[0]
        cats = [number(c) for c in plot.categories] if plot.categories is not None else []
        series = list(plot.series)
        names = [s.name or f"系列{i + 1}" for i, s in enumerate(series)]
        values = [[number(v) for v in s.values] for s in series]
        if not cats or not values:
            raise ValueError("图表无类别或系列数据")
        rows = [[cat] + [v[i] if i < len(v) else "" for v in values] for i, cat in enumerate(cats)]
        table = _render_gfm_table(["类别"] + names, rows)
        # 图表自带标题时, 把标题作为引用块压在表格前, 与 OCR 文字呈现风格一致
        return f"> **图表：{title}**\n\n{table}" if title else table
    except Exception:
        kind = str(chart.chart_type).split(" ")[0]
        note = f"*[图表：{kind}，数据无法自动提取]*"
        return f"**{title}**\n\n{note}" if title else note


def _pptx_run_text(run):
    """run 转成带行内标记的 Markdown(加粗/斜体/超链接), 首尾空白留在标记之外"""
    text = run.text
    match = re.match(r"^(\s*)(.*?)(\s*)$", text, re.DOTALL)
    if not match:
        return text
    lead, core, trail = match.groups()
    if not core:
        return text
    if run.font.bold:
        core = f"**{core}**"
    if run.font.italic:
        core = f"*{core}*"
    address = getattr(run.hyperlink, "address", None)
    if address:
        core = f"[{core}]({address})"
    return f"{lead}{core}{trail}"


def _pptx_paragraph_text(paragraph):
    """段落文本: 拼接 run 的行内格式; 域(如页码)不是 run, 天然被排除"""
    return "".join(_pptx_run_text(r) for r in paragraph.runs).strip()


def _pptx_has_own_bullet(paragraph):
    """段落自身是否带项目符号定义(文本框常见); 母版/布局继承来的符号无法据此判断"""
    from pptx.oxml.ns import qn
    try:
        pPr = paragraph._p.find(qn("a:pPr"))
        if pPr is None:
            return False
        return pPr.find(qn("a:buNone")) is None and (
            pPr.find(qn("a:buChar")) is not None or pPr.find(qn("a:buAutoNum")) is not None
        )
    except Exception:
        return False


def _join_slide_lines(lines):
    """拼一页的内容: 相邻列表项之间不留空行(紧凑列表), 其余块之间空一行"""
    out = ""
    prev_is_item = False
    for line in lines:
        is_item = re.match(r"^\s*- ", line) is not None
        if not out:
            out = line
        else:
            out += ("\n" if is_item and prev_is_item else "\n\n") + line
        prev_is_item = is_item
    return out


def _pptx_shape_lines(shape, ms, pp, images_dir, doc_stem, state, do_ocr, ocr_texts, warnings):
    """把单个形状转成 Markdown 行(可能多行/多块), 无内容返回空列表"""
    if shape.shape_type == ms.GROUP:
        lines = []
        for child in shape.shapes:  # 组合形状: 递归展开子形状
            lines += _pptx_shape_lines(child, ms, pp, images_dir, doc_stem, state, do_ocr, ocr_texts, warnings)
        return lines

    if shape.has_table:
        table = _pptx_table_to_markdown(shape.table)
        return [table] if table else []

    if shape.has_chart:
        return [_pptx_chart_to_markdown(shape.chart)]

    if shape.shape_type in (ms.PICTURE, ms.LINKED_PICTURE):
        try:
            image = shape.image
            # 图片名带页码: 与 docx 的 {文档名}_{序号} 命名区分开,
            # 同名 docx/pptx 同时转换时不会互相覆盖
            name = f"{doc_stem}_s{state['slide']}_{state['img']}.{image.ext}"
            state["img"] += 1
            rel = save_image_bytes(image.blob, images_dir, name, do_ocr, ocr_texts)
            return [f"![]({rel})"]
        except Exception as e:
            warnings.append(f"图片提取失败({type(e).__name__}): {shape.name}")
            return []

    # 无法提取内容的图形: 明确标注出来, 避免用户以为转换漏了正文
    if shape.shape_type == ms.DIAGRAM:
        return ["*[SmartArt 图形，文字无法自动提取，如需保留请截图后作为图片插入]*"]
    if shape.shape_type in (ms.EMBEDDED_OLE_OBJECT, ms.LINKED_OLE_OBJECT):
        warnings.append(f"嵌入对象未提取: {shape.name}")
        return [f"*[嵌入对象：{shape.name}]*"]
    if shape.shape_type == ms.MEDIA:
        return [f"*[媒体文件：{shape.name}]*"]

    if not shape.has_text_frame:
        return []

    ph_type = shape.placeholder_format.type if shape.is_placeholder else None
    # 页码/日期/页脚是对版面信息的重复, 不进正文
    if ph_type in (pp.SLIDE_NUMBER, pp.DATE, pp.FOOTER, pp.HEADER):
        return []

    is_title = ph_type in (pp.TITLE, pp.CENTER_TITLE, pp.VERTICAL_TITLE)
    is_subtitle = ph_type == pp.SUBTITLE
    in_placeholder = shape.is_placeholder
    lines = []
    for paragraph in shape.text_frame.paragraphs:
        text = _pptx_paragraph_text(paragraph)
        if not text:
            continue
        if is_title:
            # 标题内部换行会破坏 ATX 标题, 压成空格
            lines.append(f"## {' '.join(text.split())}")
            state["titled"] = True
        elif is_subtitle:
            lines.append(text)  # 副标题按普通段落呈现, 不加项目符号
        elif in_placeholder or paragraph.level > 0 or _pptx_has_own_bullet(paragraph):
            # 占位符正文默认继承母版的符号; 文本框只在自身带符号定义时才加"-"
            lines.append("  " * paragraph.level + f"- {text}")
        else:
            lines.append(text)
    return lines


def pptx_to_markdown(pptx_path, out_path, images_dir_name, do_ocr=True):
    """转换单个 .pptx 文件, 返回警告信息列表"""
    Presentation, ms, pp = _load_pptx()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    images_dir = out_path.parent / images_dir_name
    ocr_texts = {}
    warnings = []

    prs = Presentation(str(pptx_path))
    sections = []
    for number, slide in enumerate(prs.slides, start=1):
        # slide 页号参与图片命名, img 为页内序号; titled 记录本页是否已产出标题
        state = {"slide": number, "img": 1, "titled": False}
        lines = []
        for shape in slide.shapes:
            lines += _pptx_shape_lines(
                shape, ms, pp, images_dir, pptx_path.stem, state, do_ocr, ocr_texts, warnings
            )
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                quoted = "\n".join("> " + ln for ln in notes.splitlines())
                lines.append(f"> **备注：**\n{quoted}")
        if not lines:
            continue  # 空白页(只有页码/页脚)不产出内容
        if not state["titled"]:
            lines.insert(0, f"## 第 {number} 页")
        sections.append(_join_slide_lines(lines))

    markdown = "\n\n---\n\n".join(sections)
    markdown = insert_ocr_text(markdown, ocr_texts)
    # OCR 块自带收尾空行, 与块间分隔叠加会出现连续空行, 统一压成一个空行
    markdown = re.sub(r"\n{3,}", "\n\n", markdown)
    markdown = markdown.strip() + "\n"

    with open(out_path, "w", encoding="utf-8", newline="\n") as f:
        f.write(markdown)

    if not sections:
        warnings.append("演示文稿没有可提取的文本/图片内容")
    return warnings


def collect_input_files(inputs):
    """整理输入: 目录则递归收集 .docx/.pdf/.pptx, 文件则直接使用, 不支持的类型提示跳过"""
    files = []
    for item in inputs:
        p = Path(item)
        if p.is_dir():
            found = sorted(
                q for q in p.rglob("*")
                if q.suffix.lower() in SUPPORTED_EXTS and not q.name.startswith("~$")
            )
            if not found:
                print(f"[跳过] 目录中没有 .docx/.pdf/.pptx 文件: {p}")
            files.extend(found)
        elif p.suffix.lower() in LEGACY_EXTS:
            print(f"[跳过] 旧版 {p.suffix.lower()} 不支持, 请先另存为 {LEGACY_EXTS[p.suffix.lower()]}: {p}")
        elif p.suffix.lower() in SUPPORTED_EXTS:
            if p.exists():
                files.append(p)
            else:
                print(f"[跳过] 文件不存在: {p}")
        else:
            print(f"[跳过] 不支持的格式: {p}")
    return files


def auto_detect_files():
    """自动检测脚本所在目录(不含子目录)下的所有 .docx/.pdf/.pptx 文件"""
    script_dir = Path(__file__).resolve().parent
    print(f"未指定输入, 自动检测脚本所在目录: {script_dir}")
    files = sorted(
        p for p in script_dir.iterdir()
        if p.suffix.lower() in SUPPORTED_EXTS and not p.name.startswith("~$")
    )
    if files:
        print(f"检测到 {len(files)} 个文件: " + ", ".join(p.name for p in files) + "\n")
    return files


def main(argv=None):
    if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
        sys.stdout.reconfigure(encoding="utf-8")

    parser = argparse.ArgumentParser(
        description="Word/PDF/PPT 转 Markdown(含图片文字识别)",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__.split("用法:")[1].split("说明:")[0] if __doc__ else None,
    )
    parser.add_argument("inputs", nargs="*", help=".docx/.pdf/.pptx 文件或目录; 不填则自动检测脚本所在目录")
    parser.add_argument("-o", "--output", help="输出目录(默认与源文件同目录)")
    parser.add_argument("--images-dir", default="images", help="图片存放目录名(默认 images)")
    parser.add_argument("--no-ocr", action="store_true", help="不识别图片中的文字, 只提取图片")
    args = parser.parse_args(argv)

    files = auto_detect_files() if not args.inputs else collect_input_files(args.inputs)
    if not files:
        print("没有可转换的文件。")
        return 1

    do_ocr = not args.no_ocr
    if do_ocr and get_ocr_engine() is None:
        print("提示: 未安装 rapidocr_onnxruntime, 无法识别图片文字, 仅提取图片本身。\n"
              "      安装: pip install rapidocr_onnxruntime\n")

    ok, failed = 0, 0
    used_out_paths = set()
    converters = {
        ".docx": docx_to_markdown,
        ".pdf": pdf_to_markdown,
        ".pptx": pptx_to_markdown,
    }
    for src_path in files:
        out_dir = Path(args.output) if args.output else src_path.parent
        out_path = out_dir / (src_path.stem + ".md")
        # 同名不同格式的源文件会得到相同输出名, 追加序号避免覆盖
        base_path = out_path
        n = 2
        while out_path in used_out_paths:
            out_path = out_dir / f"{src_path.stem}({n}).md"
            n += 1
        if out_path != base_path:
            print(f"       输出重名, {src_path.name} 改为 {out_path.name}")
        used_out_paths.add(out_path)

        convert = converters.get(src_path.suffix.lower())
        if convert is None:
            print(f"[跳过] 不支持的格式: {src_path}")
            continue
        try:
            warnings = convert(src_path, out_path, args.images_dir, do_ocr)
        except Exception as e:
            failed += 1
            print(f"[失败] {src_path} -> {out_path}\n       {type(e).__name__}: {e}")
            continue
        ok += 1
        print(f"[完成] {src_path} -> {out_path}")
        # 只提示有实际影响的信息, 过滤掉常见的样式匹配噪音
        for w in warnings:
            if "style is not mapped" not in w:
                print(f"       警告: {w}")

    print(f"\n转换结束: 成功 {ok} 个, 失败 {failed} 个。")
    return 0 if failed == 0 else 2


if __name__ == "__main__":
    sys.exit(main())
